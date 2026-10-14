"""Parser routing, structure preservation and validation tests."""

from __future__ import annotations

import pytest

from rag_benchmark.parsers import choose_route, parse_document
from rag_benchmark.parsers.models import ContentType
from rag_benchmark.utils.validation import ValidationError, safe_filename


def test_choose_route_for_supported_extensions():
    assert choose_route("a.pdf") == "pdf_parser"
    assert choose_route("a.docx") == "docx_parser"
    assert choose_route("a.md") == "text_parser"
    assert choose_route("a.txt") == "text_parser"
    assert choose_route("a.html") == "html_parser"
    assert choose_route("a.csv") == "csv_parser"


def test_unsupported_extension_raises():
    with pytest.raises(ValidationError):
        choose_route("a.exe")


def test_safe_filename_blocks_traversal():
    assert ".." not in safe_filename("../../etc/passwd")
    assert "/" not in safe_filename("..\\..\\windows\\system32\\cmd.txt")


def test_markdown_headings_and_sections(markdown_file):
    doc = parse_document(markdown_file)
    headings = [b for b in doc.blocks if b.content_type == ContentType.HEADING]
    assert [h.text for h in headings] == ["Title", "Section A"]
    # Table contents inherit the full hierarchical heading path.
    table = next(b for b in doc.blocks if b.content_type == ContentType.TABLE)
    assert table.section_path == ["Title", "Section A"]


def test_markdown_table_headers_kept_with_rows(markdown_file):
    doc = parse_document(markdown_file)
    table = next(b for b in doc.blocks if b.content_type == ContentType.TABLE)
    assert table.meta["headers"] == ["Plan", "Price"]
    assert "Basic" in table.text
    assert "$49" in table.text
    assert table.table_id is not None


def test_markdown_code_block_is_code(markdown_file):
    doc = parse_document(markdown_file)
    code = [b for b in doc.blocks if b.content_type == ContentType.CODE]
    assert len(code) == 1
    assert "print('hello')" in code[0].text


def test_csv_records_keep_columns(csv_file):
    doc = parse_document(csv_file)
    records = [b for b in doc.blocks if b.content_type == ContentType.RECORD]
    assert len(records) == 2
    assert "plan: Enterprise" in records[1].text
    assert "storage_gb: 10000" in records[1].text
    assert records[0].row_number == 1
    assert doc.metadata["columns"] == ["plan", "storage_gb", "retention_days"]


def test_html_strips_scripts_and_keeps_content(html_file):
    doc = parse_document(html_file)
    full = doc.full_text
    assert "alert(1)" not in full
    assert "menu" not in full
    assert any(b.text == "Heading" for b in doc.blocks)
    assert any(b.content_type == ContentType.TABLE for b in doc.blocks)


def test_empty_text_file_warns(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")
    doc = parse_document(path)
    assert doc.block_count == 0
    assert any("No meaningful text" in w or "no extractable" in w.lower() for w in doc.warnings)


def test_corrupt_pdf_raises(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a real pdf")
    with pytest.raises(Exception):
        parse_document(path)


def test_docx_structure(tmp_path):
    from docx import Document

    document = Document()
    document.add_heading("Guide", level=1)
    document.add_paragraph("Body text.")
    document.add_heading("Steps", level=2)
    document.add_paragraph("First", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ColA"
    table.cell(0, 1).text = "ColB"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    path = tmp_path / "doc.docx"
    document.save(str(path))

    doc = parse_document(path)
    headings = [b.text for b in doc.blocks if b.content_type == ContentType.HEADING]
    assert headings == ["Guide", "Steps"]
    assert any(b.content_type == ContentType.LIST_ITEM for b in doc.blocks)
    table_block = next(b for b in doc.blocks if b.content_type == ContentType.TABLE)
    assert table_block.meta["headers"] == ["ColA", "ColB"]


def test_document_versions_and_ids_present(markdown_file):
    doc = parse_document(markdown_file)
    assert doc.source_id
    assert doc.document_version
    for block in doc.blocks:
        assert block.block_id.startswith(doc.source_id)
        assert block.parser_name
        assert block.parser_version
