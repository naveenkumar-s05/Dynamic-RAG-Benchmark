"""Generate the binary sample documents (DOCX and PDF) for the corpus.

The Markdown, TXT, HTML and CSV documents are stored as plain-text files. This
script deterministically builds the DOCX and PDF documents so the benchmark
corpus covers six formats without shipping opaque binaries by hand.

Run from the project root:

    python scripts/build_sample_docs.py
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SRC = PROJECT_ROOT / "src"
for path in (str(PROJECT_ROOT), str(SRC)):
    if path not in sys.path:
        sys.path.insert(0, path)

from config import SAMPLE_DOCS_DIR, ensure_directories  # noqa: E402


def build_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Northwind Release Notes", level=0)
    document.add_heading("Version 4.2", level=1)
    document.add_paragraph("Released 2026-03-01. This release focuses on speed and reliability.")

    document.add_heading("New Features", level=2)
    document.add_paragraph("Incremental snapshot verification", style="List Bullet")
    document.add_paragraph("Parallel restore of multiple folders", style="List Bullet")
    document.add_paragraph("Role-based access control for the console", style="List Bullet")

    document.add_heading("Fixes", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Issue"
    header[1].text = "Component"
    header[2].text = "Status"
    for issue, component, status in [
        ("E-101 on IPv6-only networks", "Agent", "Fixed"),
        ("Console slow with 10k agents", "Console", "Fixed"),
        ("Scheduled backup skipped on DST", "Scheduler", "Fixed"),
    ]:
        row = table.add_row().cells
        row[0].text = issue
        row[1].text = component
        row[2].text = status

    document.add_heading("Deprecations", level=2)
    document.add_paragraph(
        "The legacy v1 API is deprecated and will be removed in version 5.0."
    )
    document.add_paragraph(
        "Warning: configuration files without an explicit retention setting will "
        "default to 30 days after this release."
    )

    document.save(str(path))


def build_pdf(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    margin = 50
    y = margin
    line_height = 16

    def write(text: str, size: int = 11, bold: bool = False, gap: int = 1) -> None:
        nonlocal y
        fontname = "helv" if not bold else "hebo"
        page.insert_text((margin, y), text, fontsize=size, fontname=fontname)
        y += line_height * gap

    write("Northwind API Reference", size=18, bold=True, gap=1.6)
    write("Base URL: https://api.northwind.example", size=11, gap=1.4)

    write("Authentication", size=14, bold=True, gap=1.4)
    write("All requests require a bearer token in the Authorization header.")
    write("Tokens expire after 24 hours and must be refreshed.", gap=1.6)

    write("Endpoints", size=14, bold=True, gap=1.4)
    write("/v1/snapshots   GET    list snapshots", size=11)
    write("/v1/snapshots   POST   start a backup", size=11)
    write("/v1/restore     POST   restore files", size=11, gap=1.6)

    write("Rate Limits", size=14, bold=True, gap=1.4)
    write("The API allows 100 requests per minute per account.")
    write("Exceeding the limit returns HTTP 429 with a Retry-After header.", gap=1.6)

    write("Warnings", size=14, bold=True, gap=1.4)
    write("Do not call the delete endpoint for an active snapshot.")
    write("Deleting an in-progress snapshot can corrupt the recovery chain.")

    doc.save(str(path))
    doc.close()


def main() -> None:
    ensure_directories()
    SAMPLE_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    build_docx(SAMPLE_DOCS_DIR / "release_notes.docx")
    build_pdf(SAMPLE_DOCS_DIR / "api_reference.pdf")
    print(f"Wrote sample documents to {SAMPLE_DOCS_DIR}")


if __name__ == "__main__":
    main()
